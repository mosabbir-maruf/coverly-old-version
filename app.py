from flask import Flask, render_template, request, send_file
from docx import Document
import os

app = Flask(__name__, template_folder='templates')

# Path to the visitor count file
visitor_count_file = 'visitor_count.txt'

# Function to get the current visitor count
def get_visitor_count():
    if os.path.exists(visitor_count_file):
        with open(visitor_count_file, 'r') as f:
            count = int(f.read().strip())
    else:
        count = 0
    return count

# Function to update the visitor count
def increment_visitor_count():
    count = get_visitor_count() + 1
    with open(visitor_count_file, 'w') as f:
        f.write(str(count))
    return count

@app.route('/')
def home():
    # Increment the visitor count when the user visits the homepage
    visitor_count = increment_visitor_count()
    return render_template('assignment.html', visitor_count=visitor_count)  # Changed to assignment.html

@app.route('/assignment')
def assignment():
    # Increment the visitor count when the user visits the assignment page
    visitor_count = increment_visitor_count()
    return render_template('assignment.html', visitor_count=visitor_count)

@app.route('/labreport')
def labreport():
    # Increment the visitor count when the user visits the lab report page
    visitor_count = increment_visitor_count()
    return render_template('labreport.html', visitor_count=visitor_count)

@app.route('/generate_assignment', methods=['POST'])
def generate_assignment():
    # Retrieve form data
    student_id = request.form.get("student_id", "")  # Get the student ID from the form
    last_four_digits = student_id[-4:] if len(student_id) >= 4 else student_id  # Get the last 4 digits of the student ID

    # Load the document template
    doc = Document('Assignment Cover Page.docx')
    
    # Prepare the data to fill the document
    data = {
        "Here_Course_Code": request.form.get("course_code", ""),
        "No": request.form.get("assignment_no", ""),
        "Here_Course_Title": request.form.get("course_title", ""),
        "Here_TeacherName": request.form.get("teacher_name", ""),
        "teacher_designation": request.form.get("teacher_designation", ""),
        "Here_Teachers_Department_Name": request.form.get("teacher_dept", ""),
        "Here_StudentName": request.form.get("student_name", ""),
        "Here_StudentID": request.form.get("student_id", ""),
        "Here_Section": request.form.get("student_section", ""),
        "Here_DepartmentNameOfStudent": request.form.get("student_dept", ""),
        "HereDate": request.form.get("submission_date", "")
    }

    # Fill the document with the provided data
    fill_docx(doc, data)

    # Set the filename to include the last 4 digits of the student ID in the desired format
    output_filename = f'Assignment Cover Page_({last_four_digits}).docx'
    output_path = os.path.join('output', output_filename)
    
    # Save the document
    doc.save(output_path)

    # Return the document as a downloadable file
    return send_file(output_path, as_attachment=True)

@app.route('/generate_labreport', methods=['POST'])
def generate_labreport():
    # Retrieve form data
    student_id = request.form.get("student_id", "")  # Get the student ID from the form
    last_four_digits = student_id[-4:] if len(student_id) >= 4 else student_id  # Get the last 4 digits of the student ID

    # Load the document template
    doc = Document('LabReport Template.docx')
    
    # Prepare the data to fill the document
    data = {
        "Course_Code": request.form.get("course_code", ""),
        "Course_Title": request.form.get("course_title", ""),
        "HereExperimentNo": request.form.get("experiment_no", ""),
        "HereExperimentName": request.form.get("experiment_name", ""),
        "TeacherName": request.form.get("teacher_name", ""),
        "designation": request.form.get("teacher_designation", ""),
        "TeacherDeptName": request.form.get("teacher_dept", ""),
        "StudentName": request.form.get("student_name", ""),
        "StudentID": request.form.get("student_id", ""),
        "Group": request.form.get("student_section", ""),
        "DPT": request.form.get("student_dept", ""),
        "DateIs": request.form.get("submission_date", "")
    }

    # Fill the document with the provided data
    fill_docx(doc, data)

    # Set the filename to include the last 4 digits of the student ID in the desired format
    output_filename = f'Lab Report Cover Page_({last_four_digits}).docx'
    output_path = os.path.join('output', output_filename)
    
    # Save the document
    doc.save(output_path)

    # Return the document as a downloadable file
    return send_file(output_path, as_attachment=True)

def fill_docx(doc, data):
    for para in doc.paragraphs:
        for key, value in data.items():
            if key in para.text:
                for run in para.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, value)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for key, value in data.items():
                        if key in para.text:
                            for run in para.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, value)

if __name__ == "__main__":
    app.run(debug=True)

